"""
ingest.py — GST Smart Guide (.docx) -> structured chunks.jsonl

Three-pass segmentation:
  Pass 1: chapter/appendix boundaries from H-CH1 style + strict regex
  Pass 2: subsection boundaries from (a) numeric styles, (b) per-chapter
          Synopsis TOC match, (c) regex fallback
  Pass 3: noise removal (TOC paragraphs, running headers, empties)

Emits one JSON object per line to data/chunks.jsonl.
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass, field, asdict
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

# ---------------------------------------------------------------- config

TARGET_TOKENS = 400
OVERLAP_TOKENS = 60
MIN_TOKENS = 80
MAX_TOKENS = 900
TABLE_ATOMIC_ROWS = 15
TABLE_GROUP_ROWS = 8

# ---------------------------------------------------------------- regexes

CHAPTER_RE = re.compile(r"^(Chapter|Appendix)\s+(\d+)$")
NUMERIC_STYLE_RE = re.compile(r"^(\d+)\.0(\d)$")
HEADING_FALLBACK_RE = re.compile(r"^(\d+(?:\.\d+)?)[\.\s\t]+([A-Z\u201c].{3,120})$")

SECTION_RE = re.compile(r"Section\s+\d+[A-Z]?(?:\(\d+\))?(?:\([a-z]\))?")
FORM_RE = re.compile(r"FORM\s+GST\s+[A-Z]+[\s-]\d+[A-Z]?|GSTR[\s-]\d+[A-Z]?|CMP[\s-]\d+")
NOTIF_RE = re.compile(
    r"Notification\s+No\.\s*\d+/\d{4}(?:[-\u2013]\s*\w+)?(?:\s+Tax)?(?:\s*\(Rate\))?"
)
RULE_RE = re.compile(r"Rule\s+\d+[A-Z]?(?:\(\d+\))?")

DATE_PATTERNS = [
    re.compile(r"w\.e\.f\.?\s*(\d{1,2}[\.\s/-]+\w+[\.\s/-]+\d{4})", re.I),
    re.compile(r"dated\s+(\d{1,2}[\.\s/-]+\d{1,2}[\.\s/-]+\d{4})", re.I),
    re.compile(r"with effect from\s+(\d{1,2}[\.\s/-]+\w+[\.\s/-]+\d{4})", re.I),
    re.compile(r"\b((?:19|20)\d{2})\b"),  # bare year, lowest confidence
]

# ---------------------------------------------------------------- helpers


def est_tokens(text: str) -> int:
    """Cheap token estimate. Avoids a tokenizer dependency at ingest time.
    ~4 chars/token holds well for English legal prose."""
    return max(1, len(text) // 4)


CURRENCY_RE = re.compile(r"[`\u2018\u2019]\s?(?=\d)")


def normalize_currency(text: str) -> str:
    """The typesetter used a backtick as the rupee glyph in 486 places and
    the real \u20b9 in 252 others. Unnormalized, a query containing \u20b940 lakh\n    lexically misses two-thirds of the threshold passages."""
    return CURRENCY_RE.sub("\u20b9", text)


def normalize(text: str) -> str:
    """Aggressive normalization for TOC<->body heading matching.
    The typesetter used smart quotes, non-breaking spaces and inconsistent
    tabs, so naive equality fails on ~40% of true matches."""
    text = unicodedata.normalize("NFKD", text)
    text = text.replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
    text = re.sub(r"[^\w\s]", " ", text.lower())
    return re.sub(r"\s+", " ", text).strip()


def iter_block_items(doc: DocxDocument):
    """Yield paragraphs and tables in true body order.

    document.paragraphs and document.tables are separate collections, which
    loses interleaving — we need body order to attach each table to the
    section it sits inside.
    """
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def collect_header_footer_strings(doc: DocxDocument) -> set[str]:
    """The 88 docx sections leak ~80 running-header strings into the text
    stream. They contain real section titles + page numbers, so if indexed
    they outrank body text on BM25 for title-shaped queries."""
    strings = set()
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.even_page_header, section.even_page_footer,
                     section.first_page_header, section.first_page_footer):
            if part is None:
                continue
            for para in part.paragraphs:
                t = para.text.strip()
                if t:
                    strings.add(normalize(t))
    return strings


def extract_identifiers(text: str) -> dict:
    def uniq(matches):
        seen, out = set(), []
        for m in matches:
            k = re.sub(r"\s+", " ", m).strip()
            if k.lower() not in seen:
                seen.add(k.lower())
                out.append(k)
        return out

    return {
        "sections": uniq(SECTION_RE.findall(text)),
        "forms": uniq(FORM_RE.findall(text)),
        "notifications": uniq(NOTIF_RE.findall(text)),
        "rules": uniq(RULE_RE.findall(text)),
    }


def extract_effective_dates(text: str) -> list[str]:
    """Effective-date extraction drives the temporal guardrail. Confidence
    is ordered: explicit w.e.f. > dated > bare year."""
    out = []
    for pat in DATE_PATTERNS[:3]:
        out.extend(m if isinstance(m, str) else m[0] for m in pat.findall(text))
    if not out:
        years = DATE_PATTERNS[3].findall(text)
        out.extend(y for y in years if 2016 <= int(y) <= 2030)
    seen, res = set(), []
    for d in out:
        d = d.strip()
        if d not in seen:
            seen.add(d)
            res.append(d)
    return res[:6]


# ---------------------------------------------------------------- model


@dataclass
class Block:
    kind: str          # "para" | "table"
    text: str
    style: str = ""
    table_rows: list = field(default_factory=list)


@dataclass
class Chunk:
    chunk_id: str
    text: str
    raw_text: str
    chapter_num: int
    chapter_kind: str
    chapter_title: str
    section_path: list
    parent_id: str
    type: str
    identifiers: dict
    effective_dates: list
    token_count: int


# ---------------------------------------------------------------- passes


def pass1_chapters(blocks: list[Block]) -> list[dict]:
    """Chapter boundaries. Raw H-CH1 gives 109 hits, ~39 of which are body
    paragraphs misstyled by the InDesign conversion. The strict regex on
    'Chapter N' / 'Appendix N' filters those to exactly 70 real units."""
    units = []
    for i, b in enumerate(blocks):
        if b.kind != "para" or b.style != "H-CH1":
            continue
        m = CHAPTER_RE.match(b.text.strip())
        if not m:
            continue
        title = ""
        for j in range(i + 1, min(i + 4, len(blocks))):
            if blocks[j].kind == "para" and blocks[j].style == "H-CH2" and blocks[j].text.strip():
                title = blocks[j].text.strip()
                break
        units.append({
            "start": i,
            "kind": m.group(1).lower(),
            "num": int(m.group(2)),
            "title": title,
        })
    for k, u in enumerate(units):
        u["end"] = units[k + 1]["start"] if k + 1 < len(units) else len(blocks)
    return units


def chapter_toc_titles(blocks: list[Block], start: int, end: int) -> list[str]:
    """Each chapter opens with 'Synopsis' followed by toc 1 / toc 2 entries
    listing every subsection with a trailing page number. This is the
    author's own section map — we use it as ground truth for boundaries in
    the ~34 chapters that don't use numeric heading styles."""
    titles = []
    for b in blocks[start:end]:
        if b.kind == "para" and b.style.startswith("toc"):
            t = re.sub(r"\t*\d+\s*$", "", b.text).strip()
            t = re.sub(r"^\d+(\.\d+)?[\.\s\t]+", "", t).strip()
            if len(t) > 3:
                titles.append(t)
    return titles


def pass2_sections(blocks, unit, toc_titles, hf_strings, stats):
    """Split one chapter into sections. Returns list of
    {path: [...], blocks: [Block]}."""
    ch = unit["num"]
    toc_norm = {normalize(t): t for t in toc_titles}
    sections, current = [], {"path": ["_intro"], "blocks": []}
    matched_via = {"style": 0, "toc": 0, "regex": 0}
    in_synopsis = False

    for b in blocks[unit["start"]:unit["end"]]:
        if b.kind == "para":
            txt = b.text.strip()
            style = b.style

            # ---- Pass 3 (inline): noise removal
            if not txt:
                continue
            if style.startswith("toc"):
                continue
            if normalize(txt) in hf_strings:
                continue
            if txt == "Synopsis":
                in_synopsis = True
                continue
            if style in ("H-CH1", "H-CH2"):
                in_synopsis = False
                continue
            if in_synopsis and normalize(txt) in toc_norm:
                continue
            in_synopsis = False

            # ---- heading detection, three signals in priority order
            level, title = None, None

            m = NUMERIC_STYLE_RE.match(style)
            if m and int(m.group(1)) == ch and txt:
                level, title = int(m.group(2)), txt
                matched_via["style"] += 1
            elif normalize(txt) in toc_norm and len(txt) < 160:
                level, title = 1, txt
                matched_via["toc"] += 1
            else:
                fm = HEADING_FALLBACK_RE.match(txt)
                if fm and len(txt) < 120:
                    level = 2 if "." in fm.group(1) else 1
                    title = txt
                    matched_via["regex"] += 1

            if level is not None:
                if current["blocks"]:
                    sections.append(current)
                path = current["path"][:level - 1] if level > 1 else []
                current = {"path": path + [title], "blocks": []}
                continue

        current["blocks"].append(b)

    if current["blocks"]:
        sections.append(current)

    stats[ch] = {
        "toc_entries": len(toc_titles),
        "sections_found": len(sections),
        **matched_via,
    }
    return sections


# ---------------------------------------------------------------- chunking


def split_prose(text: str) -> list[str]:
    """Sentence-aware packing to TARGET_TOKENS with OVERLAP_TOKENS carry-over."""
    sents = re.split(r"(?<=[\.\?\!;])\s+(?=[A-Z\u201c(])", text)
    out, buf, buf_tok = [], [], 0
    for s in sents:
        st = est_tokens(s)
        if st > MAX_TOKENS:                      # pathological single sentence
            if buf:
                out.append(" ".join(buf)); buf, buf_tok = [], 0
            out.append(s)
            continue
        if buf_tok + st > TARGET_TOKENS and buf:
            out.append(" ".join(buf))
            carry, ct = [], 0
            for prev in reversed(buf):           # build overlap tail
                ct += est_tokens(prev)
                carry.insert(0, prev)
                if ct >= OVERLAP_TOKENS:
                    break
            buf, buf_tok = carry, ct
        buf.append(s); buf_tok += st
    if buf:
        out.append(" ".join(buf))
    # merge orphans upward
    merged = []
    for piece in out:
        if merged and est_tokens(piece) < MIN_TOKENS:
            merged[-1] = merged[-1] + " " + piece
        else:
            merged.append(piece)
    return merged


def forward_fill(body: list[list[str]]) -> list[list[str]]:
    """Some tables in this document use position instead of a real merge to
    imply grouping - one row states 'Application for Registration' in column
    0, the next several rows leave column 0 blank rather than repeating it.
    Word renders that identically to a merged cell, but there is no vMerge
    tag: it is literally blank text (verified against the source XML). If a
    row-group split later separates a blank row from the labelled row above
    it, that row loses its category with no trace it ever had one. Carrying
    the last non-empty value down through the blanks makes every row
    self-describing regardless of where a later chunk boundary falls."""
    if not body:
        return body
    last = [""] * len(body[0])
    filled = []
    for row in body:
        new_row = list(row)
        for i, cell in enumerate(new_row):
            if cell.strip():
                last[i] = cell
            else:
                new_row[i] = last[i]
        filled.append(new_row)
    return filled


def serialize_table(rows: list[list[str]]) -> list[str]:
    """<=15 rows -> one atomic chunk (preserves comparison semantics).
    Larger -> row-groups of 8 with the header row re-prepended, so every
    fragment stays independently interpretable."""
    if not rows:
        return []
    header = rows[0]
    hdr_md = "| " + " | ".join(header) + " |\n|" + "---|" * len(header)
    body = forward_fill(rows[1:])
    if len(rows) <= TABLE_ATOMIC_ROWS:
        return [hdr_md + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in body)]
    out = []
    for i in range(0, len(body), TABLE_GROUP_ROWS):
        grp = body[i:i + TABLE_GROUP_ROWS]
        out.append(hdr_md + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in grp))
    return out


def build_chunks(unit, sections) -> list[Chunk]:
    chunks = []
    kind, num = unit["kind"], unit["num"]
    ch_title = unit["title"]
    prefix = f"{kind[:3]}{num:02d}"

    for s_idx, sec in enumerate(sections):
        path = [p for p in sec["path"] if p != "_intro"]
        parent_id = f"{prefix}_s{s_idx:02d}"
        heading_line = " > ".join(
            [f"{kind.capitalize()} {num}", ch_title] + path
        )

        prose_parts = [b.text for b in sec["blocks"] if b.kind == "para"]
        tables = [b for b in sec["blocks"] if b.kind == "table"]

        pieces = []
        if prose_parts:
            joined = "\n".join(prose_parts).strip()
            if joined:
                pieces += [("prose", p) for p in split_prose(joined)]
        for t in tables:
            pieces += [("table", x) for x in serialize_table(t.table_rows)]

        for c_idx, (ptype, body) in enumerate(pieces):
            if not body.strip():
                continue
            cid = f"{parent_id}_{c_idx:03d}"
            chunks.append(Chunk(
                chunk_id=cid,
                text=f"{heading_line}\n\n{body}",
                raw_text=body,
                chapter_num=num,
                chapter_kind=kind,
                chapter_title=ch_title,
                section_path=path,
                parent_id=parent_id,
                type=ptype,
                identifiers=extract_identifiers(body),
                effective_dates=extract_effective_dates(body),
                token_count=est_tokens(body),
            ))
    return chunks


# ---------------------------------------------------------------- driver


def ingest(docx_path: str, out_path: str, report_path: str | None = None):
    doc = docx.Document(docx_path)
    hf_strings = collect_header_footer_strings(doc)

    blocks = []
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            blocks.append(Block("para", normalize_currency(item.text), item.style.name))
        else:
            rows = [[normalize_currency(c.text.strip().replace("\n", " ")) for c in r.cells]
                    for r in item.rows]
            blocks.append(Block("table", "", "", rows))

    units = pass1_chapters(blocks)
    stats, all_chunks = {}, []
    for unit in units:
        toc = chapter_toc_titles(blocks, unit["start"], unit["end"])
        secs = pass2_sections(blocks, unit, toc, hf_strings, stats)
        all_chunks.extend(build_chunks(unit, secs))

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        for c in all_chunks:
            f.write(json.dumps(asdict(c), ensure_ascii=False) + "\n")

    report = {
        "units": len(units),
        "chunks": len(all_chunks),
        "tokens_total": sum(c.token_count for c in all_chunks),
        "prose_chunks": sum(1 for c in all_chunks if c.type == "prose"),
        "table_chunks": sum(1 for c in all_chunks if c.type == "table"),
        "chunks_with_identifiers": sum(
            1 for c in all_chunks if any(c.identifiers.values())),
        "chunks_with_dates": sum(1 for c in all_chunks if c.effective_dates),
        "header_footer_strings_filtered": len(hf_strings),
        "per_chapter": stats,
    }
    if report_path:
        Path(report_path).write_text(json.dumps(report, indent=2))
    return report


if __name__ == "__main__":
    import sys
    rep = ingest(
        sys.argv[1] if len(sys.argv) > 1 else "data/manual.docx",
        sys.argv[2] if len(sys.argv) > 2 else "data/chunks.jsonl",
        "artifacts/ingest_report.json",
    )
    print(json.dumps({k: v for k, v in rep.items() if k != "per_chapter"}, indent=2))
    weak = {k: v for k, v in rep["per_chapter"].items()
            if v["sections_found"] <= 2 and v["toc_entries"] > 4}
    print(f"\nchapters needing review: {len(weak)}")
    for k, v in list(weak.items())[:12]:
        print(f"  ch{k}: toc={v['toc_entries']} found={v['sections_found']} "
              f"style={v['style']} toc_m={v['toc']} regex={v['regex']}")

